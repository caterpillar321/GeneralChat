"""파일별 텍스트 추출 (페이지 단위 메타 보존)."""

from __future__ import annotations

from pathlib import Path

from pypdf import PdfReader

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def extract_pdf_pages(path: Path) -> list[tuple[int, str]]:
    reader = PdfReader(str(path))
    pages: list[tuple[int, str]] = []
    for i, page in enumerate(reader.pages):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        pages.append((i + 1, text))
    return pages


def extract_docx(path: Path) -> list[tuple[int, str]]:
    """DOCX 는 페이지 메타 없음 — 단일 가상 페이지로 통합."""
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    parts: list[str] = []

    for para in doc.paragraphs:
        text = (para.text or "").strip()
        if text:
            parts.append(text)

    # 표 셀 텍스트
    for table in doc.tables:
        for row in table.rows:
            cells = [(cell.text or "").strip() for cell in row.cells]
            cells = [c for c in cells if c]
            if cells:
                parts.append(" | ".join(cells))

    return [(1, "\n\n".join(parts))]


def extract_pptx(path: Path) -> list[tuple[int, str]]:
    """슬라이드별 텍스트 + 노트. 슬라이드 번호를 page_num 으로 매핑."""
    from pptx import Presentation

    prs = Presentation(str(path))
    pages: list[tuple[int, str]] = []
    for i, slide in enumerate(prs.slides):
        chunks: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text or "" for run in para.runs).strip()
                    if line:
                        chunks.append(line)
            elif getattr(shape, "has_table", False):
                table = shape.table
                for row in table.rows:
                    cells = [(c.text or "").strip() for c in row.cells]
                    cells = [c for c in cells if c]
                    if cells:
                        chunks.append(" | ".join(cells))
        # 슬라이드 노트
        if slide.has_notes_slide:
            note = (slide.notes_slide.notes_text_frame.text or "").strip()
            if note:
                chunks.append(f"[발표자 노트]\n{note}")
        text = "\n".join(chunks).strip()
        if text:
            pages.append((i + 1, text))
    return pages


def extract_text(path: str | Path, mime: str | None = None) -> list[tuple[int, str]]:
    """파일 형식에 따라 라우팅. (page_num, text) 리스트 반환."""
    p = Path(path)
    suffix = p.suffix.lower()

    if suffix == ".pdf" or mime == "application/pdf":
        return extract_pdf_pages(p)
    if suffix == ".docx" or mime == DOCX_MIME:
        return extract_docx(p)
    if suffix == ".pptx" or mime == PPTX_MIME:
        return extract_pptx(p)
    if suffix in (".txt", ".md", ".markdown") or (mime or "").startswith("text/"):
        return [(1, p.read_text(encoding="utf-8", errors="replace"))]

    raise ValueError(f"지원 안 되는 파일 형식: {suffix} ({mime})")
